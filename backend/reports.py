"""PDF and DOCX report generators for compliance reports."""
import io
from datetime import datetime
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from docx import Document
from docx.shared import Pt


def _fmt(v):
    return str(v) if v not in (None, "", "null") else "—"


def build_pdf_report(scan: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18*mm, bottomMargin=18*mm, leftMargin=18*mm, rightMargin=18*mm)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], textColor=colors.HexColor("#0F172A"), fontSize=18, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=colors.HexColor("#1E293B"), fontSize=13, spaceAfter=4)
    normal = styles["BodyText"]
    small = ParagraphStyle("small", parent=styles["BodyText"], fontSize=9, textColor=colors.HexColor("#475569"))

    story = []
    story.append(Paragraph("Legal Metrology Compliance Report", h1))
    story.append(Paragraph("Packaged Commodities Rules, 2011", small))
    story.append(Spacer(1, 8))

    verdict = scan.get("compliance", {}).get("verdict", "N/A")
    score = scan.get("compliance", {}).get("score", 0)
    color_map = {"COMPLIANT": "#10B981", "PARTIALLY_COMPLIANT": "#F59E0B", "NON_COMPLIANT": "#EF4444"}
    vcolor = color_map.get(verdict, "#475569")

    meta_data = [
        ["Report ID", scan.get("id", "")],
        ["Scanned By", scan.get("scanned_by_email", "")],
        ["Scan Date", scan.get("created_at", "")],
        ["Verdict", verdict.replace("_", " ")],
        ["Compliance Score", f"{score}%"],
    ]
    t = Table(meta_data, colWidths=[45*mm, 130*mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1F5F9")),
        ("TEXTCOLOR", (1, 3), (1, 3), colors.HexColor(vcolor)),
        ("FONTNAME", (1, 3), (1, 3), "Helvetica-Bold"),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E8F0")),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(t)
    story.append(Spacer(1, 12))

    story.append(Paragraph("Extracted Declarations", h2))
    ext = scan.get("extracted", {})
    fields = [
        ("Product Name", ext.get("product_name")),
        ("Manufacturer", ext.get("manufacturer_name")),
        ("Manufacturer Address", ext.get("manufacturer_address")),
        ("Net Quantity", ext.get("net_quantity")),
        ("MRP", ext.get("mrp")),
        ("Mfg Date", ext.get("mfg_date")),
        ("Expiry Date", ext.get("expiry_date")),
        ("Consumer Care", ext.get("consumer_care")),
        ("Country of Origin", ext.get("country_of_origin")),
        ("Ingredients", ext.get("ingredients")),
        ("FSSAI Number", ext.get("fssai_number")),
        ("Batch Number", ext.get("batch_number")),
    ]
    data = [["Field", "Value"]] + [[k, _fmt(v)] for k, v in fields]
    tbl = Table(data, colWidths=[55*mm, 120*mm])
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E8F0")),
        ("PADDING", (0, 0), (-1, -1), 5),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(tbl)
    story.append(Spacer(1, 12))

    story.append(Paragraph("Rule-by-Rule Compliance Checks", h2))
    checks = scan.get("compliance", {}).get("checks", [])
    cdata = [["Rule", "Declaration", "Status", "Note"]]
    for c in checks:
        cdata.append([c["rule"], c["label"], c["status"].upper(), c.get("note", "")])
    ctbl = Table(cdata, colWidths=[25*mm, 70*mm, 22*mm, 58*mm])
    ctbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E8F0")),
        ("PADDING", (0, 0), (-1, -1), 5),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    for i, c in enumerate(checks, start=1):
        st = c["status"]
        col = "#10B981" if st == "pass" else ("#F59E0B" if st == "warning" else "#EF4444")
        ctbl.setStyle(TableStyle([("TEXTCOLOR", (2, i), (2, i), colors.HexColor(col)), ("FONTNAME", (2, i), (2, i), "Helvetica-Bold")]))
    story.append(ctbl)
    story.append(Spacer(1, 10))

    violations = scan.get("compliance", {}).get("violations", [])
    if violations:
        story.append(Paragraph("Violations & Warnings", h2))
        for v in violations:
            story.append(Paragraph(f"• {v}", normal))
            story.append(Spacer(1, 4))

    story.append(Spacer(1, 20))
    story.append(Paragraph(f"Digitally generated on {datetime.utcnow().strftime('%d %b %Y %H:%M UTC')} — Legal Metrology Compliance Engine", small))

    doc.build(story)
    return buf.getvalue()


def build_docx_report(scan: dict) -> bytes:
    d = Document()
    d.add_heading("Legal Metrology Compliance Report", level=0)
    d.add_paragraph("Packaged Commodities Rules, 2011").italic = True

    verdict = scan.get("compliance", {}).get("verdict", "N/A")
    score = scan.get("compliance", {}).get("score", 0)

    d.add_heading("Report Metadata", level=1)
    tbl = d.add_table(rows=5, cols=2)
    tbl.style = "Light Grid Accent 1"
    meta = [
        ("Report ID", scan.get("id", "")),
        ("Scanned By", scan.get("scanned_by_email", "")),
        ("Scan Date", str(scan.get("created_at", ""))),
        ("Verdict", verdict.replace("_", " ")),
        ("Compliance Score", f"{score}%"),
    ]
    for i, (k, v) in enumerate(meta):
        tbl.rows[i].cells[0].text = k
        tbl.rows[i].cells[1].text = str(v)

    d.add_heading("Extracted Declarations", level=1)
    ext = scan.get("extracted", {})
    fields = [
        ("Product Name", ext.get("product_name")),
        ("Manufacturer", ext.get("manufacturer_name")),
        ("Manufacturer Address", ext.get("manufacturer_address")),
        ("Net Quantity", ext.get("net_quantity")),
        ("MRP", ext.get("mrp")),
        ("Mfg Date", ext.get("mfg_date")),
        ("Expiry Date", ext.get("expiry_date")),
        ("Consumer Care", ext.get("consumer_care")),
        ("Country of Origin", ext.get("country_of_origin")),
        ("Ingredients", ext.get("ingredients")),
        ("FSSAI Number", ext.get("fssai_number")),
        ("Batch Number", ext.get("batch_number")),
    ]
    ft = d.add_table(rows=len(fields) + 1, cols=2)
    ft.style = "Light Grid Accent 1"
    ft.rows[0].cells[0].text = "Field"
    ft.rows[0].cells[1].text = "Value"
    for i, (k, v) in enumerate(fields, start=1):
        ft.rows[i].cells[0].text = k
        ft.rows[i].cells[1].text = _fmt(v)

    d.add_heading("Rule-by-Rule Compliance Checks", level=1)
    checks = scan.get("compliance", {}).get("checks", [])
    ct = d.add_table(rows=len(checks) + 1, cols=4)
    ct.style = "Light Grid Accent 1"
    hdr = ["Rule", "Declaration", "Status", "Note"]
    for i, h in enumerate(hdr):
        ct.rows[0].cells[i].text = h
    for i, c in enumerate(checks, start=1):
        ct.rows[i].cells[0].text = c["rule"]
        ct.rows[i].cells[1].text = c["label"]
        ct.rows[i].cells[2].text = c["status"].upper()
        ct.rows[i].cells[3].text = c.get("note", "")

    violations = scan.get("compliance", {}).get("violations", [])
    if violations:
        d.add_heading("Violations & Warnings", level=1)
        for v in violations:
            d.add_paragraph(v, style="List Bullet")

    d.add_paragraph()
    footer = d.add_paragraph(f"Digitally generated on {datetime.utcnow().strftime('%d %b %Y %H:%M UTC')} — Legal Metrology Compliance Engine")
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(9)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
